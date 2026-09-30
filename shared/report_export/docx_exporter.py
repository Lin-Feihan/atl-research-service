from pathlib import Path
import re

import pypandoc

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = Path(__file__).parent

DEFAULT_REFERENCE_DOC = (
    BASE_DIR
    / "templates"
    / "default_report.docx"
)


# =====================================================
# General helpers
# =====================================================

def extract_document_title(markdown_text):
    """
    Extract the first Markdown heading and use it
    as the DOCX document title.
    """

    for line in markdown_text.splitlines():
        match = re.match(
            r"^\s*#{1,6}\s+(.+?)\s*$",
            line,
        )

        if match:
            return match.group(1).strip()

    return "Research Report"


# =====================================================
# DOCX XML helpers
# =====================================================

def set_cell_shading(cell, fill):
    """
    Set table cell background color.
    """

    tc_pr = cell._tc.get_or_add_tcPr()

    # Remove existing shading so our value wins
    existing = tc_pr.findall(
        qn("w:shd")
    )

    for node in existing:
        tc_pr.remove(node)

    shd = OxmlElement(
        "w:shd"
    )

    shd.set(
        qn("w:val"),
        "clear",
    )

    shd.set(
        qn("w:color"),
        "auto",
    )

    shd.set(
        qn("w:fill"),
        fill,
    )

    tc_pr.append(
        shd
    )


def set_cell_margins(
    cell,
    top=120,
    start=120,
    bottom=120,
    end=120,
):
    """
    Set internal table-cell padding.

    Values are in twips.
    """

    tc_pr = (
        cell
        ._tc
        .get_or_add_tcPr()
    )

    tc_mar = (
        tc_pr
        .first_child_found_in(
            "w:tcMar"
        )
    )

    if tc_mar is None:
        tc_mar = OxmlElement(
            "w:tcMar"
        )

        tc_pr.append(
            tc_mar
        )

    for margin, value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ]:

        node = tc_mar.find(
            qn(
                f"w:{margin}"
            )
        )

        if node is None:
            node = OxmlElement(
                f"w:{margin}"
            )

            tc_mar.append(
                node
            )

        node.set(
            qn("w:w"),
            str(value),
        )

        node.set(
            qn("w:type"),
            "dxa",
        )


def set_table_borders(table):
    """
    Apply light gray table borders.
    """

    tbl_pr = (
        table
        ._tbl
        .tblPr
    )

    borders = (
        tbl_pr
        .first_child_found_in(
            "w:tblBorders"
        )
    )

    if borders is None:
        borders = OxmlElement(
            "w:tblBorders"
        )

        tbl_pr.append(
            borders
        )

    for border_name in [
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV",
    ]:

        border = borders.find(
            qn(
                f"w:{border_name}"
            )
        )

        if border is None:
            border = OxmlElement(
                f"w:{border_name}"
            )

            borders.append(
                border
            )

        border.set(
            qn("w:val"),
            "single",
        )

        border.set(
            qn("w:sz"),
            "4",
        )

        border.set(
            qn("w:space"),
            "0",
        )

        border.set(
            qn("w:color"),
            "D1D5DB",
        )


def set_repeat_table_header(row):
    """
    Repeat the first table row when a table
    continues onto another page.
    """

    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement(
        "w:tblHeader"
    )

    tbl_header.set(
        qn("w:val"),
        "true",
    )

    tr_pr.append(
        tbl_header
    )


# =====================================================
# Hyperlink styling
# =====================================================

def configure_hyperlink_style(doc):
    """
    Make hyperlinks professional but still identifiable.

    Dark blue-gray + underline rather than Word's
    bright default blue.
    """

    styles = doc.styles

    try:
        hyperlink = styles[
            "Hyperlink"
        ]

    except KeyError:
        hyperlink = styles.add_style(
            "Hyperlink",
            WD_STYLE_TYPE.CHARACTER,
        )

    hyperlink.font.name = (
        "Arial"
    )

    hyperlink.font.size = (
        Pt(10.5)
    )

    hyperlink.font.color.rgb = (
        RGBColor(
            55,
            65,
            81,
        )
    )

    hyperlink.font.underline = True


# =====================================================
# Paragraph formatting
# =====================================================

def format_paragraphs(doc):
    """
    Improve paragraph pagination behavior.
    """

    for paragraph in doc.paragraphs:

        style_name = (
            paragraph.style.name
            if paragraph.style
            else ""
        )

        # Reduce widows/orphans in normal body text
        if style_name == "Normal":
            paragraph.paragraph_format.widow_control = True

        # Keep headings with the paragraph that follows
        if style_name.startswith(
            "Heading"
        ):
            paragraph.paragraph_format.keep_with_next = True

        # Avoid paragraphs being split unnecessarily
        if (
            style_name == "Normal"
            and len(paragraph.text) < 120
        ):
            paragraph.paragraph_format.keep_together = True


# =====================================================
# Table helpers
# =====================================================

def is_wide_table(table):
    """
    >= 5 columns are treated as wide tables.

    This version does not force landscape sections yet.
    Wide tables simply receive slightly smaller text.
    """

    if not table.rows:
        return False

    return (
        len(
            table.rows[0].cells
        )
        >= 5
    )


def average_first_column_length(table):
    """
    Estimate whether the first column contains short IDs
    such as F-001, R-002, S-003, etc.
    """

    values = []

    for row in table.rows[1:]:
        if not row.cells:
            continue

        value = (
            row.cells[0]
            .text
            .strip()
        )

        if value:
            values.append(
                len(value)
            )

    if not values:
        return None

    return (
        sum(values)
        / len(values)
    )


def optimize_first_column(table):
    """
    Narrow short identifier columns such as:

    Finding ID
    F-001
    F-002

    This is deliberately conservative.
    """

    if not table.rows:
        return

    column_count = len(
        table.rows[0].cells
    )

    if column_count < 2:
        return

    avg_length = (
        average_first_column_length(
            table
        )
    )

    if avg_length is None:
        return

    # Only narrow genuinely short identifier columns
    if avg_length <= 12:

        target_width = Inches(
            0.9
        )

        for row in table.rows:
            row.cells[0].width = (
                target_width
            )


# =====================================================
# Table formatting
# =====================================================

def format_tables(doc):
    """
    Apply shared professional formatting
    to every table.
    """

    for table in doc.tables:

        # Keep table structure visible and predictable
        table.style = (
            "Table Grid"
        )

        set_table_borders(
            table
        )

        wide = is_wide_table(
            table
        )

        # Allow Word to calculate the rest of the widths
        table.autofit = True

        for row in table.rows:

            for cell in row.cells:

                cell.vertical_alignment = (
                    WD_CELL_VERTICAL_ALIGNMENT.CENTER
                )

                set_cell_margins(
                    cell,
                    top=120,
                    start=120,
                    bottom=120,
                    end=120,
                )

                for paragraph in cell.paragraphs:

                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT
                    )

                    paragraph.paragraph_format.space_before = (
                        Pt(0)
                    )

                    paragraph.paragraph_format.space_after = (
                        Pt(2)
                    )

                    paragraph.paragraph_format.line_spacing = (
                        1.0
                    )

                    for run in paragraph.runs:

                        run.font.name = (
                            "Arial"
                        )

                        # Slightly smaller for wide tables
                        run.font.size = (
                            Pt(8.5)
                            if wide
                            else Pt(9)
                        )

        # ---------------------------------------------
        # Header row
        # ---------------------------------------------

        if table.rows:

            header_row = (
                table.rows[0]
            )

            set_repeat_table_header(
                header_row
            )

            for cell in header_row.cells:

                set_cell_shading(
                    cell,
                    "E9EDF2",
                )

                for paragraph in cell.paragraphs:

                    paragraph.paragraph_format.space_after = (
                        Pt(2)
                    )

                    for run in paragraph.runs:

                        run.bold = True

                        run.font.name = (
                            "Arial"
                        )

                        run.font.size = (
                            Pt(8.5)
                            if wide
                            else Pt(9)
                        )

        # ---------------------------------------------
        # Column-width heuristic
        # ---------------------------------------------

        optimize_first_column(
            table
        )


# =====================================================
# Document metadata
# =====================================================

def format_document_metadata(
    doc,
    document_title,
):
    """
    Set clean DOCX metadata.
    """

    properties = (
        doc.core_properties
    )

    properties.title = (
        document_title
    )

    properties.subject = (
        "Research Report"
    )

    properties.author = (
        "Agentic Trading Lab"
    )

    properties.keywords = (
        "research, due diligence, "
        "deep research, agent"
    )

    properties.comments = (
        "Generated through the "
        "Agentic Trading Lab research workflow."
    )


# =====================================================
# Output validation
# =====================================================

def validate_output(doc):
    """
    Check for common Markdown residue.

    This does not stop report generation.
    It only prints warnings.
    """

    markdown_link_pattern = re.compile(
        r"\[[^\]]+\]\([^)]+\)"
    )

    heading_pattern = re.compile(
        r"^\s*#{1,6}\s+"
    )

    warnings = []

    for paragraph in doc.paragraphs:

        text = (
            paragraph.text
            .strip()
        )

        if not text:
            continue

        if markdown_link_pattern.search(
            text
        ):
            warnings.append(
                text
            )

        elif heading_pattern.search(
            text
        ):
            warnings.append(
                text
            )

        elif text == "---":
            warnings.append(
                text
            )

    if warnings:

        print(
            "Warning: possible Markdown "
            "residue detected:"
        )

        for warning in warnings[:10]:

            print(
                f"  - {warning[:120]}"
            )


# =====================================================
# Main Markdown -> DOCX exporter
# =====================================================

def markdown_to_docx(
    markdown_text,
    output_path,
    reference_doc=None,
):
    """
    Convert Markdown to a professionally formatted DOCX.

    This is the shared DOCX renderer used by
    all research agents.
    """

    output_path = Path(
        output_path
    )

    # Ensure output folder exists
    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    # ---------------------------------------------
    # Reference template
    # ---------------------------------------------

    if reference_doc is None:

        reference_doc = (
            DEFAULT_REFERENCE_DOC
        )

    reference_doc = Path(
        reference_doc
    )

    extra_args = [
        "--standalone",
    ]

    if reference_doc.exists():

        extra_args.append(
            f"--reference-doc={reference_doc}"
        )

    # ---------------------------------------------
    # Step 1:
    # Markdown -> DOCX using Pandoc
    # ---------------------------------------------

    try:

        pypandoc.convert_text(
            markdown_text,
            to="docx",
            format="gfm",
            outputfile=str(
                output_path
            ),
            extra_args=extra_args,
        )

    except Exception as exc:

        raise RuntimeError(
            "Pandoc DOCX conversion failed: "
            f"{exc}"
        ) from exc

    if not output_path.exists():

        raise RuntimeError(
            "DOCX file was not created."
        )

    # ---------------------------------------------
    # Step 2:
    # DOCX post-processing
    # ---------------------------------------------

    doc = Document(
        output_path
    )

    document_title = (
        extract_document_title(
            markdown_text
        )
    )

    configure_hyperlink_style(
        doc
    )

    format_document_metadata(
        doc,
        document_title,
    )

    format_paragraphs(
        doc
    )

    format_tables(
        doc
    )

    validate_output(
        doc
    )

    # ---------------------------------------------
    # Step 3:
    # Save final DOCX
    # ---------------------------------------------

    doc.save(
        output_path
    )

    return output_path
