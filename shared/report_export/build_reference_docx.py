from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


BASE_DIR = Path(__file__).parent
TEMPLATE_PATH = BASE_DIR / "templates" / "default_report.docx"


def set_font(
    style,
    name,
    size,
    bold=False,
    color="000000",
):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)

    # Make Word use the same font for all character sets
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()

    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")

    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "

    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_1)
    run._r.append(instr_text)
    run._r.append(fld_char_2)

    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)


def build_reference_docx():
    doc = Document(TEMPLATE_PATH)

    # --------------------------------------------------
    # Page layout
    # --------------------------------------------------
    for section in doc.sections:
        section.top_margin = Inches(0.80)
        section.bottom_margin = Inches(0.80)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)

        # Clean footer + page number
        footer = section.footer

        paragraph = footer.paragraphs[0]
        paragraph.clear()

        add_page_number(paragraph)

    # --------------------------------------------------
    # Normal body text
    # --------------------------------------------------
    normal = doc.styles["Normal"]

    set_font(
        normal,
        "Arial",
        10.5,
        color="111111",
    )

    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    # --------------------------------------------------
    # Title
    # --------------------------------------------------
    if "Title" in doc.styles:
        title = doc.styles["Title"]

        set_font(
            title,
            "Arial",
            22,
            bold=True,
            color="1F2937",
        )

        title.paragraph_format.space_before = Pt(0)
        title.paragraph_format.space_after = Pt(12)

    # --------------------------------------------------
    # Heading 1
    # --------------------------------------------------
    heading1 = doc.styles["Heading 1"]

    set_font(
        heading1,
        "Arial",
        16,
        bold=True,
        color="1F2937",
    )

    heading1.paragraph_format.space_before = Pt(16)
    heading1.paragraph_format.space_after = Pt(7)
    heading1.paragraph_format.keep_with_next = True

    # --------------------------------------------------
    # Heading 2
    # --------------------------------------------------
    heading2 = doc.styles["Heading 2"]

    set_font(
        heading2,
        "Arial",
        12.5,
        bold=True,
        color="374151",
    )

    heading2.paragraph_format.space_before = Pt(12)
    heading2.paragraph_format.space_after = Pt(5)
    heading2.paragraph_format.keep_with_next = True

    # --------------------------------------------------
    # Heading 3
    # --------------------------------------------------
    heading3 = doc.styles["Heading 3"]

    set_font(
        heading3,
        "Arial",
        10.5,
        bold=True,
        color="4B5563",
    )

    heading3.paragraph_format.space_before = Pt(8)
    heading3.paragraph_format.space_after = Pt(3)
    heading3.paragraph_format.keep_with_next = True

    doc.save(TEMPLATE_PATH)

    print(
        f"Reference DOCX updated: {TEMPLATE_PATH}"
    )


if __name__ == "__main__":
    build_reference_docx()